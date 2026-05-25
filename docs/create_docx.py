from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    """셀 배경색 설정"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)

def create_table(doc, headers, rows, header_color='D9D9D9'):
    """테이블 생성"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'

    # 헤더
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_shading(header_cells[i], header_color)

    # 데이터
    for row_idx, row_data in enumerate(rows):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_data in enumerate(row_data):
            row_cells[col_idx].text = str(cell_data) if cell_data else ''

    return table

doc = Document()

# 제목
title = doc.add_heading('3. 구현단계', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_heading('3.1 프로그램 코드', level=1)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

# 제·개정 이력
doc.add_heading('제·개정 이력', level=2)
create_table(doc,
    ['날짜', '버전', '작성자', '승인자', '내용'],
    [
        ['2026-04-21', '1.0', '개발팀', '', '최초 작성'],
        ['', '', '', '', ''],
        ['', '', '', '', ''],
    ])

doc.add_paragraph()

# I1. 프로그램 코드
doc.add_heading('I1. 프로그램 코드', level=2)
create_table(doc,
    ['항목', '내용'],
    [
        ['시스템명', 'AutoFlow RPA Platform'],
        ['서브시스템명', 'AI 기반 자동화 워크플로우 생성 시스템'],
        ['단계명', '구현'],
        ['작성일자', '2026-04-21'],
        ['버전', '1.0'],
    ])

doc.add_paragraph()

# 1. 프로그램 목록
doc.add_heading('1. 프로그램 목록', level=2)

# 1.1 Frontend - 애플리케이션 페이지
doc.add_heading('1.1 Frontend - 애플리케이션 페이지 (app/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['페이지 라우팅', 'FE-APP-001', 'app/layout.tsx', '루트 레이아웃, 전역 스타일 및 프로바이더 설정'],
        ['페이지 라우팅', 'FE-APP-002', 'app/page.tsx', '메인 페이지, 애플리케이션 진입점'],
    ])

doc.add_paragraph()

# 1.2 Frontend - 인증 컴포넌트
doc.add_heading('1.2 Frontend - 인증 컴포넌트 (components/auth/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['사용자 인증', 'FE-AUTH-001', 'components/auth/auth-page.tsx', '로그인/회원가입 페이지 컴포넌트'],
    ])

doc.add_paragraph()

# 1.3 Frontend - 워크플로우 블록 컴포넌트
doc.add_heading('1.3 Frontend - 워크플로우 블록 컴포넌트 (components/blocks/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['AI 워크플로우 생성', 'FE-BLK-001', 'components/blocks/ai-create-dialog.tsx', 'AI 자연어 입력 다이얼로그'],
        ['AI 워크플로우 생성', 'FE-BLK-002', 'components/blocks/ai-preview-modal.tsx', 'AI 생성 결과 미리보기 모달'],
        ['블록 관리', 'FE-BLK-003', 'components/blocks/block-palette.tsx', '120+ 블록 팔레트 UI'],
        ['블록 관리', 'FE-BLK-004', 'components/blocks/draggable-block.tsx', '드래그 앤 드롭 블록 컴포넌트'],
        ['실행 관리', 'FE-BLK-005', 'components/blocks/execution-panel.tsx', '워크플로우 실행 패널'],
        ['변수 관리', 'FE-BLK-006', 'components/blocks/variables-panel.tsx', '변수 설정 패널'],
        ['비디오 분석', 'FE-BLK-007', 'components/blocks/video-upload-modal.tsx', '비디오 업로드 및 분석 모달'],
        ['블록 관리', 'FE-BLK-008', 'components/blocks/workspace-block.tsx', '개별 블록 렌더링 컴포넌트'],
        ['캔버스 관리', 'FE-BLK-009', 'components/blocks/workspace-canvas.tsx', '워크플로우 캔버스 컴포넌트'],
    ])

doc.add_paragraph()

# 1.4 Frontend - 빌더 컴포넌트
doc.add_heading('1.4 Frontend - 빌더 컴포넌트 (components/builder/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['워크스페이스 관리', 'FE-BLD-001', 'components/builder/multi-workspace.tsx', '멀티 워크스페이스 관리 컴포넌트'],
    ])

doc.add_paragraph()

# 1.5 Frontend - 채팅 컴포넌트
doc.add_heading('1.5 Frontend - 채팅 컴포넌트 (components/chat/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['AI 채팅', 'FE-CHT-001', 'components/chat/chat-page.tsx', 'AI 어시스턴트 채팅 페이지'],
    ])

doc.add_paragraph()

# 1.6 Frontend - 마켓플레이스 컴포넌트
doc.add_heading('1.6 Frontend - 마켓플레이스 컴포넌트 (components/marketplace/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['마켓플레이스', 'FE-MKT-001', 'components/marketplace/marketplace-card.tsx', '상품 카드 컴포넌트'],
        ['마켓플레이스', 'FE-MKT-002', 'components/marketplace/marketplace-detail.tsx', '상품 상세 정보 컴포넌트'],
        ['마켓플레이스', 'FE-MKT-003', 'components/marketplace/marketplace-detail-page.tsx', '상품 상세 페이지'],
        ['마켓플레이스', 'FE-MKT-004', 'components/marketplace/marketplace-page.tsx', '마켓플레이스 메인 페이지'],
        ['구매 관리', 'FE-MKT-005', 'components/marketplace/purchase-modal.tsx', '상품 구매 모달'],
        ['판매 관리', 'FE-MKT-006', 'components/marketplace/sell-modal.tsx', '상품 판매 등록 모달'],
    ])

doc.add_paragraph()

# 1.7 Frontend - 마이페이지 컴포넌트
doc.add_heading('1.7 Frontend - 마이페이지 컴포넌트 (components/mypage/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['사용자 관리', 'FE-MYP-001', 'components/mypage/my-page.tsx', '마이페이지 (프로필, 구매내역, 판매내역)'],
    ])

doc.add_paragraph()

# 1.8 Frontend - UI 기본 컴포넌트
doc.add_heading('1.8 Frontend - UI 기본 컴포넌트 (components/ui/)', level=3)

ui_components = [
    ['UI 컴포넌트', 'FE-UI-001', 'components/ui/accordion.tsx', '아코디언 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-002', 'components/ui/alert-dialog.tsx', '경고 다이얼로그'],
    ['UI 컴포넌트', 'FE-UI-003', 'components/ui/alert.tsx', '알림 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-004', 'components/ui/aspect-ratio.tsx', '비율 컨테이너'],
    ['UI 컴포넌트', 'FE-UI-005', 'components/ui/avatar.tsx', '아바타 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-006', 'components/ui/badge.tsx', '배지 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-007', 'components/ui/breadcrumb.tsx', '브레드크럼 네비게이션'],
    ['UI 컴포넌트', 'FE-UI-008', 'components/ui/button-group.tsx', '버튼 그룹'],
    ['UI 컴포넌트', 'FE-UI-009', 'components/ui/button.tsx', '버튼 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-010', 'components/ui/calendar.tsx', '캘린더 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-011', 'components/ui/card.tsx', '카드 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-012', 'components/ui/carousel.tsx', '캐러셀 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-013', 'components/ui/chart.tsx', '차트 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-014', 'components/ui/checkbox.tsx', '체크박스 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-015', 'components/ui/collapsible.tsx', '접을 수 있는 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-016', 'components/ui/command.tsx', '커맨드 팔레트'],
    ['UI 컴포넌트', 'FE-UI-017', 'components/ui/context-menu.tsx', '컨텍스트 메뉴'],
    ['UI 컴포넌트', 'FE-UI-018', 'components/ui/dialog.tsx', '다이얼로그 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-019', 'components/ui/drawer.tsx', '드로어 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-020', 'components/ui/dropdown-menu.tsx', '드롭다운 메뉴'],
    ['UI 컴포넌트', 'FE-UI-021', 'components/ui/empty.tsx', '빈 상태 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-022', 'components/ui/field.tsx', '폼 필드 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-023', 'components/ui/form.tsx', '폼 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-024', 'components/ui/hover-card.tsx', '호버 카드'],
    ['UI 컴포넌트', 'FE-UI-025', 'components/ui/input-group.tsx', '입력 그룹'],
    ['UI 컴포넌트', 'FE-UI-026', 'components/ui/input-otp.tsx', 'OTP 입력'],
    ['UI 컴포넌트', 'FE-UI-027', 'components/ui/input.tsx', '입력 필드'],
    ['UI 컴포넌트', 'FE-UI-028', 'components/ui/item.tsx', '아이템 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-029', 'components/ui/kbd.tsx', '키보드 단축키 표시'],
    ['UI 컴포넌트', 'FE-UI-030', 'components/ui/label.tsx', '레이블 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-031', 'components/ui/menubar.tsx', '메뉴바 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-032', 'components/ui/navigation-menu.tsx', '네비게이션 메뉴'],
    ['UI 컴포넌트', 'FE-UI-033', 'components/ui/pagination.tsx', '페이지네이션'],
    ['UI 컴포넌트', 'FE-UI-034', 'components/ui/popover.tsx', '팝오버 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-035', 'components/ui/progress.tsx', '프로그레스 바'],
    ['UI 컴포넌트', 'FE-UI-036', 'components/ui/radio-group.tsx', '라디오 버튼 그룹'],
    ['UI 컴포넌트', 'FE-UI-037', 'components/ui/resizable.tsx', '리사이즈 가능 패널'],
    ['UI 컴포넌트', 'FE-UI-038', 'components/ui/scroll-area.tsx', '스크롤 영역'],
    ['UI 컴포넌트', 'FE-UI-039', 'components/ui/select.tsx', '셀렉트 박스'],
    ['UI 컴포넌트', 'FE-UI-040', 'components/ui/separator.tsx', '구분선'],
    ['UI 컴포넌트', 'FE-UI-041', 'components/ui/sheet.tsx', '시트 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-042', 'components/ui/sidebar.tsx', '사이드바 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-043', 'components/ui/skeleton.tsx', '스켈레톤 로딩'],
    ['UI 컴포넌트', 'FE-UI-044', 'components/ui/slider.tsx', '슬라이더 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-045', 'components/ui/sonner.tsx', 'Sonner 토스트'],
    ['UI 컴포넌트', 'FE-UI-046', 'components/ui/spinner.tsx', '로딩 스피너'],
    ['UI 컴포넌트', 'FE-UI-047', 'components/ui/switch.tsx', '스위치 토글'],
    ['UI 컴포넌트', 'FE-UI-048', 'components/ui/table.tsx', '테이블 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-049', 'components/ui/tabs.tsx', '탭 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-050', 'components/ui/textarea.tsx', '텍스트 영역'],
    ['UI 컴포넌트', 'FE-UI-051', 'components/ui/toast.tsx', '토스트 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-052', 'components/ui/toaster.tsx', '토스터 프로바이더'],
    ['UI 컴포넌트', 'FE-UI-053', 'components/ui/toggle-group.tsx', '토글 그룹'],
    ['UI 컴포넌트', 'FE-UI-054', 'components/ui/toggle.tsx', '토글 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-055', 'components/ui/tooltip.tsx', '툴팁 컴포넌트'],
    ['UI 컴포넌트', 'FE-UI-056', 'components/ui/use-mobile.tsx', '모바일 감지 훅'],
    ['UI 컴포넌트', 'FE-UI-057', 'components/ui/use-toast.ts', '토스트 훅'],
]
create_table(doc, ['서브시스템명', '프로그램파일 ID', '파일명', '비고'], ui_components)

doc.add_paragraph()

# 1.9 Frontend - 루트 컴포넌트
doc.add_heading('1.9 Frontend - 루트 컴포넌트 (components/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['자동화 빌더', 'FE-CMP-001', 'components/automation-builder.tsx', '메인 자동화 빌더 컴포넌트'],
        ['테마 관리', 'FE-CMP-002', 'components/theme-provider.tsx', '다크/라이트 테마 프로바이더'],
    ])

doc.add_paragraph()

# 1.10 Frontend - 라이브러리
doc.add_heading('1.10 Frontend - 라이브러리 (lib/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['API 통신', 'FE-LIB-001', 'lib/api.ts', 'API 클라이언트 (40+ 메서드)'],
        ['유틸리티', 'FE-LIB-002', 'lib/utils.ts', '공통 유틸리티 함수'],
    ])

doc.add_paragraph()

# 1.11 Frontend - 타입 정의
doc.add_heading('1.11 Frontend - 타입 정의 (types/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['타입 정의', 'FE-TYP-001', 'types/blocks.ts', '블록 타입 정의 (120+ 블록, 10개 카테고리)'],
    ])

doc.add_paragraph()

# 1.12 Frontend - 커스텀 훅
doc.add_heading('1.12 Frontend - 커스텀 훅 (hooks/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['훅', 'FE-HOK-001', 'hooks/use-mobile.ts', '모바일 환경 감지 훅'],
        ['훅', 'FE-HOK-002', 'hooks/use-toast.ts', '토스트 알림 훅'],
    ])

doc.add_page_break()

# 1.13 Backend - API 라우터
doc.add_heading('1.13 Backend - API 라우터 (backend/app/api/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['API 초기화', 'BE-API-001', 'backend/app/api/__init__.py', 'API 라우터 초기화'],
        ['AI 서비스', 'BE-API-002', 'backend/app/api/ai.py', 'AI 분석 및 워크플로우 생성 API'],
        ['인증', 'BE-API-003', 'backend/app/api/auth.py', 'JWT 기반 인증 API'],
        ['채팅', 'BE-API-004', 'backend/app/api/chat.py', 'AI 채팅 API'],
        ['실행', 'BE-API-005', 'backend/app/api/execution.py', '코드 생성 및 워크플로우 실행 API'],
        ['지식베이스', 'BE-API-006', 'backend/app/api/knowledge.py', '지식베이스 검색 API (458개 프로젝트)'],
        ['마켓플레이스', 'BE-API-007', 'backend/app/api/marketplace.py', '마켓플레이스 CRUD API'],
        ['구매', 'BE-API-008', 'backend/app/api/purchases.py', '구매 처리 API'],
        ['사용자', 'BE-API-009', 'backend/app/api/users.py', '사용자 관리 API'],
        ['비디오', 'BE-API-010', 'backend/app/api/videos.py', '비디오 업로드 및 분석 API'],
        ['WebSocket', 'BE-API-011', 'backend/app/api/websocket.py', '실시간 통신 WebSocket API'],
        ['워크스페이스', 'BE-API-012', 'backend/app/api/workspaces.py', '워크플로우 CRUD API'],
    ])

doc.add_paragraph()

# 1.14 Backend - 데이터 모델
doc.add_heading('1.14 Backend - 데이터 모델 (backend/app/models/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['모델 초기화', 'BE-MDL-001', 'backend/app/models/__init__.py', '모델 초기화 및 export'],
        ['대화 모델', 'BE-MDL-002', 'backend/app/models/conversation.py', '대화 세션 모델'],
        ['마켓플레이스', 'BE-MDL-003', 'backend/app/models/marketplace_item.py', '마켓플레이스 아이템 모델'],
        ['메시지 모델', 'BE-MDL-004', 'backend/app/models/message.py', '채팅 메시지 모델'],
        ['구매 모델', 'BE-MDL-005', 'backend/app/models/purchase.py', '구매 내역 모델'],
        ['리뷰 모델', 'BE-MDL-006', 'backend/app/models/review.py', '상품 리뷰 모델'],
        ['템플릿 모델', 'BE-MDL-007', 'backend/app/models/template.py', '워크플로우 템플릿 모델'],
        ['사용자 모델', 'BE-MDL-008', 'backend/app/models/user.py', '사용자 계정 모델'],
        ['비디오 모델', 'BE-MDL-009', 'backend/app/models/video.py', '업로드 비디오 모델'],
        ['워크스페이스', 'BE-MDL-010', 'backend/app/models/workspace.py', '워크스페이스 모델'],
    ])

doc.add_paragraph()

# 1.15 Backend - Pydantic 스키마
doc.add_heading('1.15 Backend - Pydantic 스키마 (backend/app/schemas/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['스키마 초기화', 'BE-SCH-001', 'backend/app/schemas/__init__.py', '스키마 초기화'],
        ['AI 스키마', 'BE-SCH-002', 'backend/app/schemas/ai.py', 'AI 요청/응답 스키마'],
        ['인증 스키마', 'BE-SCH-003', 'backend/app/schemas/auth.py', '인증 관련 스키마'],
        ['채팅 스키마', 'BE-SCH-004', 'backend/app/schemas/chat.py', '채팅 메시지 스키마'],
        ['마켓플레이스', 'BE-SCH-005', 'backend/app/schemas/marketplace.py', '마켓플레이스 스키마'],
        ['구매 스키마', 'BE-SCH-006', 'backend/app/schemas/purchase.py', '구매 관련 스키마'],
        ['템플릿 스키마', 'BE-SCH-007', 'backend/app/schemas/template.py', '템플릿 스키마'],
        ['비디오 스키마', 'BE-SCH-008', 'backend/app/schemas/video.py', '비디오 관련 스키마'],
        ['워크스페이스', 'BE-SCH-009', 'backend/app/schemas/workspace.py', '워크스페이스 스키마'],
    ])

doc.add_paragraph()

# 1.16 Backend - 비즈니스 로직 서비스
doc.add_heading('1.16 Backend - 비즈니스 로직 서비스 (backend/app/services/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['코드 생성', 'BE-SVC-001', 'backend/app/services/code_generator.py', '블록 → Python 코드 변환'],
        ['비디오 처리', 'BE-SVC-002', 'backend/app/services/video_processor.py', 'FFmpeg + Claude Vision 비디오 분석'],
        ['웹 분석', 'BE-SVC-003', 'backend/app/services/web_analyzer.py', '웹 페이지 분석 서비스'],
    ])

doc.add_paragraph()

# 1.17 Backend - 지식베이스 서비스
doc.add_heading('1.17 Backend - 지식베이스 서비스 (backend/app/services/knowledge_base/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['지식베이스 초기화', 'BE-KB-001', 'backend/app/services/knowledge_base/__init__.py', '지식베이스 모듈 초기화'],
        ['임베딩', 'BE-KB-002', 'backend/app/services/knowledge_base/embeddings.py', 'OpenAI 임베딩 생성'],
        ['인덱서', 'BE-KB-003', 'backend/app/services/knowledge_base/indexer.py', '프로젝트 인덱싱 (458개)'],
        ['지식베이스', 'BE-KB-004', 'backend/app/services/knowledge_base/knowledge_base.py', 'ChromaDB 지식베이스 관리'],
        ['검색', 'BE-KB-005', 'backend/app/services/knowledge_base/search.py', '하이브리드 검색 (키워드+시맨틱)'],
    ])

doc.add_paragraph()

# 1.18 Backend - 유틸리티
doc.add_heading('1.18 Backend - 유틸리티 (backend/app/utils/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['유틸 초기화', 'BE-UTL-001', 'backend/app/utils/__init__.py', '유틸리티 모듈 초기화'],
        ['보안', 'BE-UTL-002', 'backend/app/utils/security.py', '비밀번호 해싱, JWT 토큰 관리'],
    ])

doc.add_paragraph()

# 1.19 Backend - 애플리케이션 코어
doc.add_heading('1.19 Backend - 애플리케이션 코어 (backend/app/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['앱 초기화', 'BE-APP-001', 'backend/app/__init__.py', '앱 패키지 초기화'],
        ['설정', 'BE-APP-002', 'backend/app/config.py', '환경 설정 (API 키, DB 등)'],
        ['데이터베이스', 'BE-APP-003', 'backend/app/database.py', 'SQLAlchemy 비동기 DB 연결'],
        ['메인', 'BE-APP-004', 'backend/app/main.py', 'FastAPI 애플리케이션 진입점'],
    ])

doc.add_paragraph()

# 1.20 Backend - 루트 스크립트
doc.add_heading('1.20 Backend - 루트 스크립트 (backend/)', level=3)
create_table(doc,
    ['서브시스템명', '프로그램파일 ID', '파일명', '비고'],
    [
        ['실행 스크립트', 'BE-RUN-001', 'backend/run.py', '개발 서버 실행 스크립트'],
        ['지식베이스 빌드', 'BE-RUN-002', 'backend/build_knowledge_base.py', '지식베이스 구축 스크립트'],
        ['마이그레이션', 'BE-RUN-003', 'backend/migrate_add_notification_columns.py', 'DB 마이그레이션 스크립트'],
    ])

doc.add_page_break()

# 2. 프로그램 소스코드
doc.add_heading('2. 프로그램 소스코드', level=2)
doc.add_paragraph('소스코드는 별도 제출합니다.')

# 2.1 소스코드 저장소 정보
doc.add_heading('2.1 소스코드 저장소 정보', level=3)
create_table(doc,
    ['항목', '내용'],
    [
        ['저장소 위치', 'C:\\Users\\user\\Desktop\\RPA'],
        ['VCS 유형', 'Git'],
        ['메인 브랜치', 'main'],
    ])

doc.add_paragraph()

# 2.2 소스코드 구조
doc.add_heading('2.2 소스코드 구조', level=3)
structure = """RPA/
├── app/                          # Next.js 페이지 (2개 파일)
├── components/                   # React 컴포넌트
│   ├── auth/                     # 인증 (1개 파일)
│   ├── blocks/                   # 워크플로우 블록 (9개 파일)
│   ├── builder/                  # 빌더 (1개 파일)
│   ├── chat/                     # 채팅 (1개 파일)
│   ├── marketplace/              # 마켓플레이스 (6개 파일)
│   ├── mypage/                   # 마이페이지 (1개 파일)
│   └── ui/                       # UI 컴포넌트 (57개 파일)
├── lib/                          # 유틸리티 (2개 파일)
├── types/                        # 타입 정의 (1개 파일)
├── hooks/                        # 커스텀 훅 (2개 파일)
└── backend/                      # FastAPI 백엔드
    └── app/
        ├── api/                  # API 라우터 (12개 파일)
        ├── models/               # ORM 모델 (10개 파일)
        ├── schemas/              # Pydantic 스키마 (9개 파일)
        ├── services/             # 비즈니스 로직 (8개 파일)
        └── utils/                # 유틸리티 (2개 파일)"""

p = doc.add_paragraph()
run = p.add_run(structure)
run.font.name = 'Consolas'
run.font.size = Pt(9)

doc.add_paragraph()

# 2.3 파일 통계 요약
doc.add_heading('2.3 파일 통계 요약', level=3)
create_table(doc,
    ['구분', '서브시스템', '파일 수'],
    [
        ['Frontend', '페이지 (app/)', '2'],
        ['Frontend', '인증 컴포넌트', '1'],
        ['Frontend', '블록 컴포넌트', '9'],
        ['Frontend', '빌더 컴포넌트', '1'],
        ['Frontend', '채팅 컴포넌트', '1'],
        ['Frontend', '마켓플레이스 컴포넌트', '6'],
        ['Frontend', '마이페이지 컴포넌트', '1'],
        ['Frontend', 'UI 컴포넌트', '57'],
        ['Frontend', '루트 컴포넌트', '2'],
        ['Frontend', '라이브러리', '2'],
        ['Frontend', '타입 정의', '1'],
        ['Frontend', '커스텀 훅', '2'],
        ['Frontend 소계', '', '85'],
        ['Backend', 'API 라우터', '12'],
        ['Backend', '데이터 모델', '10'],
        ['Backend', 'Pydantic 스키마', '9'],
        ['Backend', '비즈니스 로직', '3'],
        ['Backend', '지식베이스', '5'],
        ['Backend', '유틸리티', '2'],
        ['Backend', '애플리케이션 코어', '4'],
        ['Backend', '루트 스크립트', '3'],
        ['Backend 소계', '', '48'],
        ['총계', '', '133'],
    ])

doc.add_paragraph()

# 작성 목적
doc.add_heading('작성 목적', level=2)
doc.add_paragraph('설계 명세서에서 기술한 프로그램 코드에 대한 물리적인 형상과 코드 관리를 위한 명세를 기술한다.')

# 작성 방법
doc.add_heading('작성 방법', level=2)
doc.add_paragraph('프로그램 코드의 프로그램 목록을 서브시스템별로 분류하여 기술하였다.')

# 항목 설명
doc.add_heading('항목 설명', level=2)
doc.add_heading('프로그램 목록', level=3)
doc.add_paragraph('• 서브시스템명: 본 프로그램파일이 관련되는 서브시스템명을 기입')
doc.add_paragraph('• 프로그램파일 ID: 프로그램 파일의 고유 ID (FE-: Frontend, BE-: Backend)')
doc.add_paragraph('• 프로그램파일명: 프로그램 파일의 상대 경로 및 파일명')
doc.add_paragraph('• 비고: 해당 파일의 주요 기능 및 역할 설명')

doc.add_paragraph()
doc.add_paragraph('─' * 40)
p = doc.add_paragraph('문서 끝')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 저장
doc.save('C:\\Users\\user\\Desktop\\RPA\\docs\\3.1_프로그램코드_산출물.docx')
print('Word 문서가 생성되었습니다: docs/3.1_프로그램코드_산출물.docx')
